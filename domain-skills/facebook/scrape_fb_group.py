import argparse
import json
import time
from datetime import datetime, timezone, timedelta
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from browser_harness.helpers import (
    cdp,
    drain_events,
    list_tabs,
    switch_tab,
    new_tab,
    js,
)

def find_stories_in_json(obj, stories_list):
    if isinstance(obj, dict):
        if obj.get("__typename") == "Story" or ("comet_sections" in obj and "post_id" in obj):
            stories_list.append(obj)
            return
        if "edges" in obj and isinstance(obj["edges"], list):
            for edge in obj["edges"]:
                if isinstance(edge, dict) and "node" in edge:
                    node = edge["node"]
                    if isinstance(node, dict) and (node.get("__typename") == "Story" or "comet_sections" in node):
                        stories_list.append(node)
                        continue
        for v in obj.values():
            find_stories_in_json(v, stories_list)
    elif isinstance(obj, list):
        for item in obj:
            find_stories_in_json(item, stories_list)

def extract_story_info(story, group_url):
    post_id = story.get("post_id") or story.get("id")
    
    author = ""
    actors = story.get("actors", [])
    if actors and isinstance(actors, list) and isinstance(actors[0], dict):
        author = actors[0].get("name", "")
    
    ts = story.get("creation_time")
    if not ts and story.get("tracking"):
        try:
            tr = json.loads(story["tracking"]) if isinstance(story["tracking"], str) else story["tracking"]
            for p_info in tr.get("page_insights", {}).values():
                ts = p_info.get("post_context", {}).get("publish_time")
                if ts:
                    break
        except Exception:
            pass
    date_str = ""
    if ts:
        vn_tz = timezone(timedelta(hours=7))
        date_str = datetime.fromtimestamp(ts, tz=vn_tz).strftime("%Y-%m-%d %H:%M:%S")

    content = ""
    msg_obj = story.get("comet_sections", {}).get("content", {}).get("story", {}).get("comet_sections", {}).get("message", {})
    rich_msg = msg_obj.get("rich_message", [])
    if rich_msg:
        texts = [b.get("text", "") for b in rich_msg if b.get("text")]
        content = "\n".join(texts).strip()
    if not content:
        content = msg_obj.get("story", {}).get("message", {}).get("text", "")
    if not content:
        content = story.get("comet_sections", {}).get("feedback", {}).get("story", {}).get("story_ufi_container", {}).get("story", {}).get("message", {}).get("text", "")
    if not content:
        content = story.get("message", {}).get("text", "")
    if not content and story.get("attached_story"):
        att_story = story["attached_story"]
        att_msg = att_story.get("comet_sections", {}).get("content", {}).get("story", {}).get("comet_sections", {}).get("message", {})
        if att_msg.get("rich_message"):
            content = "\n".join([b.get("text", "") for b in att_msg["rich_message"] if b.get("text")])
        elif att_msg.get("story", {}).get("message", {}).get("text"):
            content = att_msg["story"]["message"]["text"]

    images = []
    def find_imgs(att):
        styles = att.get("styles", {}).get("attachment", {})
        media_list = [styles.get("media", {}), att.get("media", {})]
        sub = styles.get("all_subattachments", {}).get("nodes", [])
        for s in sub:
            media_list.append(s.get("media", {}))
        for m in media_list:
            if not isinstance(m, dict):
                continue
            for k in ["photo_image", "image", "large_share_image"]:
                uri = m.get(k, {}).get("uri")
                if uri and uri not in images:
                    images.append(uri)
    for att in story.get("attachments", []):
        find_imgs(att)
    if not images and story.get("attached_story"):
        for att in story["attached_story"].get("attachments", []):
            find_imgs(att)

    likes = 0
    comments = 0
    fb_target = story.get("comet_sections", {}).get("feedback", {}).get("story", {}).get("story_ufi_container", {}).get("story", {}).get("feedback_context", {}).get("feedback_target_with_context", {})
    renderer = fb_target.get("comet_ufi_summary_and_actions_renderer", {}).get("feedback", {})
    actions = renderer.get("adaptive_ufi_action_renderers", [])
    for a in actions:
        typename = a.get("__typename")
        fb = a.get("feedback", {})
        if typename == "UFIStoryReactActionRenderer":
            likes = fb.get("reaction_count", {}).get("count", 0)
        elif typename == "UFICommentActionRenderer":
            comments = fb.get("comment_rendering_instance", {}).get("comments", {}).get("total_count", 0)

    if not likes:
        top_r = renderer.get("top_reactions", {}).get("edges", [])
        likes = sum(e.get("reaction_count", 0) for e in top_r if isinstance(e, dict))

    permalink = story.get("permalink_url") or story.get("url") or ""
    if not permalink and post_id:
        clean_group = group_url.rstrip("/")
        permalink = f"{clean_group}/posts/{post_id}/"

    return {
        "id": post_id,
        "author": author,
        "date": date_str,
        "content": content,
        "images": "\n".join(images),
        "likes": likes,
        "comments": comments,
        "permalink": permalink
    }

def export_to_excel(posts, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facebook Group Posts"

    headers = ["STT", "Ngày đăng", "Tác giả", "Nội dung", "Link ảnh", "Lượt Like", "Lượt Comments", "Link bài viết"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1877F2", end_color="1877F2", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, p in enumerate(posts, 1):
        row = [
            idx,
            p["date"],
            p["author"],
            p["content"],
            p["images"],
            p["likes"],
            p["comments"],
            p["permalink"]
        ]
        ws.append(row)
        current_row = idx + 1
        for col_num in range(1, len(row) + 1):
            c = ws.cell(row=current_row, column=col_num)
            c.border = thin_border
            if col_num in [1, 2, 6, 7]:
                c.alignment = Alignment(horizontal="center", vertical="top")
            else:
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    col_widths = {1: 8, 2: 20, 3: 20, 4: 60, 5: 35, 6: 12, 7: 15, 8: 35}
    for col_num, width in col_widths.items():
        col_letter = get_column_letter(col_num)
        ws.column_dimensions[col_letter].width = width

    resolved_path = Path(output_path).resolve()
    resolved_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(resolved_path)
    return resolved_path

def scrape_group(group_url, limit, output_path):
    tabs = list_tabs()
    target_id = None
    clean_url = group_url.rstrip("/")
    for t in tabs:
        t_url = t.get("url", "").rstrip("/")
        if clean_url in t_url:
            target_id = t["targetId"]
            break

    if target_id:
        switch_tab(target_id, activate=True)
    else:
        new_tab(group_url)

    cdp("Network.enable")
    time.sleep(0.5)
    drain_events()

    collected_stories = []
    seen_ids = set()
    stagnant_count = 0
    prev_count = 0

    max_scrolls = max(50, limit * 3)
    for scroll_idx in range(max_scrolls):
        js("window.scrollBy(0, 2200)")
        time.sleep(1.8)
        events = drain_events()
        for ev in events:
            if ev.get("method") == "Network.responseReceived":
                r = ev.get("params", {}).get("response", {})
                if "graphql" in r.get("url", ""):
                    req_id = ev.get("params", {}).get("requestId")
                    try:
                        body = cdp("Network.getResponseBody", requestId=req_id).get("body", "")
                        stories_in_body = []
                        for line in body.split("\n"):
                            line = line.strip()
                            if not line:
                                continue
                            try:
                                data = json.loads(line)
                                find_stories_in_json(data, stories_in_body)
                            except Exception:
                                pass
                        for s in stories_in_body:
                            info = extract_story_info(s, group_url)
                            sid = info["id"]
                            if sid and sid not in seen_ids and (info["content"] or info["images"]):
                                seen_ids.add(sid)
                                collected_stories.append(info)
                                if len(collected_stories) % 10 == 0 or len(collected_stories) == limit:
                                    print(f"[{len(collected_stories)}/{limit}] posts collected")
                    except Exception:
                        pass
        if len(collected_stories) >= limit:
            break
        if len(collected_stories) == prev_count:
            stagnant_count += 1
            if stagnant_count >= 12:
                js("window.scrollTo(0, document.body.scrollHeight)")
                time.sleep(3)
            if stagnant_count >= 25:
                break
        else:
            stagnant_count = 0
            prev_count = len(collected_stories)

    saved_file = export_to_excel(collected_stories[:limit], output_path)
    print(f"Scraped {len(collected_stories[:limit])} posts successfully -> {saved_file}")

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--url", required=True)
    parser.add_argument("--limit", type=int, default=50)
    parser.add_argument("--output", default="fb_posts.xlsx")
    args = parser.parse_args()

    scrape_group(args.url, args.limit, args.output)

if __name__ == "__main__":
    main()
